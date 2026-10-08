import copy
import json
import os
import tempfile
import unittest
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch
from PySide6.QtWidgets import QApplication, QMainWindow
from PySide6.QtCore import Qt, QCoreApplication, QEvent
from app.services.storage import Store, timestamp
from app.services.accounts import Accounts, ConnectionUnavailable
from app.services.contributions import ActiveSession, Participation, WeeklyReports, aggregate, export_pdf, export_csv, report_html
from app.services.meetings import Meetings, import_transcript
from app.ui.participation import MeetingEditor
from tests import test_research_workspace as fixtures

QT_APP = None
USER = '9f020e44-7f50-4aef-a716-655a647c2498'
OTHER = '7f47b79b-96cd-4d64-b473-f731f5548541'


class ParticipationTests(unittest.TestCase):
    def test_idle_resume_focus_logout_monotonic(self):
        clock=[0.0]
        session=ActiveSession(USER,str(uuid.uuid4()),lambda:clock[0],datetime(2026,10,5,tzinfo=timezone.utc))
        clock[0]=300;session.touch()
        clock[0]=2100;session.tick()
        self.assertEqual(session.row['active_seconds'],1200)  # 5 min activity + 15 min cutoff
        session.touch();clock[0]=2160;session.focus(False)
        clock[0]=6000;session.tick();self.assertEqual(session.row['active_seconds'],1260)
        session.focus(True);clock[0]=6060
        self.assertEqual(session.close()['active_seconds'],1320)
        clock[0]=7000;self.assertEqual(session.close()['ended_at'],session.at(6060))
        self.assertEqual(len(session.row['active_intervals']),3)

    def test_weekly_counts_owner_isolation_overlap_and_boundaries(self):
        begin=datetime(2026,10,5,tzinfo=timezone.utc);end=begin+timedelta(days=7)
        events=[]
        for i,kind in enumerate(['LOGIN','LOGIN','APP_STARTED','FILE_ADDED','FILE_OPENED','SAMPLE_VIEWED','EXPERIMENT_UPDATED','IMAGE_VIEWED','REPORT_OPENED','MEETING_NOTES_EDITED','TRANSCRIPT_EDITED','ACTION_ITEM_COMPLETED']):
            details={'source':'legacy','file_category':'Data'}
            if kind=='FILE_ADDED':details={'attribution':'app_user_action','file_category':'Images','image_category':'Original'}
            events.append(dict(id=i,client_event_id=str(uuid.uuid4()),user_id=USER,event_type=kind,created_at=(begin+timedelta(hours=i)).isoformat(),entity_type='sample' if kind=='SAMPLE_VIEWED' else 'experiment' if kind=='EXPERIMENT_UPDATED' else 'file',entity_id='S-holy' if kind=='SAMPLE_VIEWED' else 'E-1',entity_name='Holy GDL',details=details))
        events += [dict(events[0]),dict(events[0],id=50,user_id=OTHER,client_event_id=str(uuid.uuid4())),dict(events[0],id=51,created_at=end.isoformat(),client_event_id=str(uuid.uuid4()))]
        sessions=[dict(id='a',user_id=USER,started_at=begin.isoformat(),last_active_at=(begin+timedelta(minutes=60)).isoformat(),ended_at=None,active_intervals=[[begin.isoformat(),(begin+timedelta(minutes=60)).isoformat()]]),
                  dict(id='b',user_id=USER,started_at=(begin+timedelta(minutes=30)).isoformat(),last_active_at=(begin+timedelta(minutes=90)).isoformat(),active_intervals=[[(begin+timedelta(minutes=30)).isoformat(),(begin+timedelta(minutes=90)).isoformat()]])]
        meetings=[dict(id='m1',title='Imaging review',meeting_date='2026-10-05'),dict(id='m2',title='Advisor',meeting_date='2026-10-06')]
        attendance=[dict(user_id=USER,meeting_id='m1',attended=True,attendance_status='Present'),dict(user_id=USER,meeting_id='m2',attended=False,attendance_status='Absent'),dict(user_id=OTHER,meeting_id='m2',attended=True,attendance_status='Present')]
        report=aggregate(dict(id=USER,display_name='Researcher A'),begin,end,events,sessions,attendance,meetings)
        metrics=report['metrics']
        self.assertEqual((metrics['Logins'],metrics['Sessions'],metrics['Active time']),(2,2,'1 h 30 min'))
        self.assertEqual(metrics['Attendance'],'1 / 2 (50%)')
        self.assertEqual(metrics['Files added (known attribution)'],1)
        self.assertEqual(metrics['Unique samples worked with'],1)
        self.assertEqual(metrics['Unique experiments worked with'],1)
        self.assertEqual(metrics['Data files opened'],1)
        self.assertEqual(metrics['Original images added'],1)
        other=aggregate(dict(id=OTHER,display_name='Researcher B'),begin,end,events,sessions,attendance,meetings)
        self.assertEqual(other['metrics']['Logins'],1);self.assertEqual(other['metrics']['Files opened'],0)
        self.assertNotIn('Researcher B',report_html(report))


class MeetingTests(unittest.TestCase):
    def setUp(self):
        global QT_APP
        QT_APP=QApplication.instance() or QApplication([])
        self.fixture=fixtures.WorkspaceTests();self.fixture.setUp()
        self.store=self.fixture.store;self.cloud={};self.events={};self.sessions={};self.online=True
        self.profile=dict(id=USER,email='member@example.invalid',display_name='Researcher A',role='user',active=True)
        def transport(method,path,data):
            if not self.online:raise ConnectionUnavailable('Offline')
            if path.startswith('/rest/v1/activity_events') and method=='POST':
                for e in data:self.events.setdefault(e['client_event_id'],copy.deepcopy(e))
            elif path.startswith('/rest/v1/user_sessions') and method=='POST':self.sessions[data['id']]=copy.deepcopy(data)
            elif path=='/rest/v1/rpc/hub_save_meeting':
                row=copy.deepcopy(data['record']);old=self.cloud.get(row['id'])
                if old and old['client_mutation_id']==row['client_mutation_id']:return old
                if old and old['revision']!=data['expected_revision']:return {'conflict':True}
                row['revision']=data['expected_revision']+1;self.cloud[row['id']]=row;return copy.deepcopy(row)
            elif path=='/rest/v1/rpc/hub_team_roster':return [self.profile,dict(id=OTHER,display_name='Researcher B',active=True)]
            elif path.startswith('/rest/v1/meetings?'):return list(copy.deepcopy(self.cloud).values())
            elif path.startswith('/rest/v1/meeting_attendees?'):return [dict(a,meeting_id=m['id']) for m in self.cloud.values() for a in m['attendees']]
            elif path.startswith('/rest/v1/meeting_actions?'):return [dict(a,meeting_id=m['id']) for m in self.cloud.values() for a in m['actions']]
            elif path.startswith('/rest/v1/activity_events?') and method=='GET':return list(copy.deepcopy(self.events).values())
            elif path.startswith('/rest/v1/user_sessions?') and method=='GET':return list(copy.deepcopy(self.sessions).values())
            return []
        self.accounts=Accounts(self.store,transport);self.accounts.profile=self.profile;self.accounts.session={'access_token':'test'}
        self.store.accounts=self.accounts;self.service=Meetings(self.accounts,self.fixture.catalog)

    def tearDown(self):
        self.accounts.executor.shutdown(wait=True)
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
        QApplication.processEvents()
        import gc
        gc.collect()
        self.fixture.tearDown()

    def test_meeting_restart_attendance_revision_retry_conflict_and_owner(self):
        row=self.service.new();row['title']='Imaging review';row['content']['notes']='Decisions and next steps'
        row['content']['transcript']='Matthew: Verify Holy GDL.'
        row['attendees']=[dict(person_key=USER,user_id=USER,guest_name='',attended=True,attendance_status='Present',notes=''),dict(person_key=OTHER,user_id=OTHER,guest_name='',attended=False,attendance_status='Absent',notes='')]
        row['actions']=[dict(id=str(uuid.uuid4()),title='Verify sample family',assigned_to=USER,due_date='2026-10-16',completed=False)]
        self.online=False;self.service.save(row)
        with self.assertRaises(ValueError):self.service.sync()
        reopened=Meetings(self.accounts,self.fixture.catalog)
        self.assertEqual(reopened.list('Holy')[0]['content']['notes'],row['content']['notes'])
        self.online=True;self.assertEqual(reopened.sync(),'Meetings synchronized')
        shared=reopened.get(row['id']);self.assertEqual(shared['revision'],1)
        self.assertFalse(shared['attendees'][1]['attended'])
        changed=copy.deepcopy(shared);changed['content']['transcript']='Edited transcript';changed['actions'][0]['completed']=True
        reopened.save(changed);reopened.sync();reopened.sync()
        self.assertEqual(reopened.get(row['id'])['revision'],2)
        self.assertEqual(reopened.get(row['id'])['content']['transcript_history'][-1]['text'],'Matthew: Verify Holy GDL.')
        stale=copy.deepcopy(shared);stale['content']['notes']='Stale edits';reopened.save(stale)
        self.assertIn('Conflict',reopened.sync());self.assertEqual(reopened.get(row['id'])['content']['notes'],'Stale edits')
        self.assertEqual(self.cloud[row['id']]['content']['notes'],'Decisions and next steps')
        other=Meetings(self.accounts,self.fixture.catalog);other.user=OTHER
        with self.assertRaises(ValueError):other.save(shared)

    def test_session_offline_absolute_retry_and_event_dedup(self):
        p=Participation(self.accounts);p.session.row['active_seconds']=0
        self.online=False
        with self.assertRaises(ConnectionUnavailable):p.flush()
        self.assertTrue(p.pending)
        self.online=True;p.flush();p.flush();self.assertEqual(len(self.sessions),1)
        self.accounts.offline=True;self.accounts.event('FILE_OPENED','file','f','report.pdf')
        snapshot=copy.deepcopy(self.accounts.pending);self.accounts.offline=False;self.accounts.flush()
        self.accounts.pending=snapshot;self.accounts.flush()
        self.assertEqual(len(self.events),1)

    def test_docx_and_extractable_pdf_transcript_import(self):
        from docx import Document
        from reportlab.pdfgen import canvas
        doc=Document();doc.add_paragraph('Speaker: Review Holy GDL.');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Decision';table.cell(0,1).text='Verify pore map'
        path=self.fixture.current/'transcript.docx';doc.save(path)
        text=import_transcript(path);self.assertIn('Holy GDL',text);self.assertIn('Verify pore map',text)
        path=self.fixture.current/'transcript.pdf';pdf=canvas.Canvas(str(path));pdf.drawString(40,700,'Speaker: Verify the sample.');pdf.save()
        self.assertIn('Verify the sample',import_transcript(path))

    def test_file_attribution_external_scan_and_app_copy(self):
        self.accounts.offline=True
        (self.fixture.current/'external.csv').write_text('value\n1',encoding='utf-8')
        self.fixture.catalog.refresh()
        self.assertFalse(any(e['event_type']=='FILE_ADDED' for e in self.accounts.pending))
        destination=self.fixture.catalog.import_current(self.fixture.old)
        added=[e for e in self.accounts.pending if e['event_type']=='FILE_ADDED']
        self.assertEqual(len(added),1);self.assertEqual(added[0]['user_id'],USER)
        self.assertEqual(added[0]['details']['attribution'],'app_user_action')
        self.assertTrue(destination.is_file())

    def test_ui_attendance_recording_transcript_notes_actions_and_portable_paths(self):
        from types import SimpleNamespace
        self.accounts.offline=True
        audio=self.fixture.current/'review.wav';audio.write_bytes(b'RIFF'+b'0'*128)
        transcript=self.fixture.current/'review.md';transcript.write_text('Andrew: Holy GDL review.',encoding='utf-8')
        ref=self.service.file_ref(audio);self.assertNotIn(str(self.fixture.root),json.dumps(ref))
        self.assertEqual(self.service.file_path(ref),audio)
        with self.assertRaises(ValueError):self.service.file_path(dict(ref,relative_path='../outside.wav'))
        parent=QMainWindow();parent.window=SimpleNamespace(research_workspace=SimpleNamespace(repository=self.fixture.repo))
        row=self.service.new();row['title']='Review';row['content']['files']['recording']=ref
        editor=MeetingEditor(self.service,row,[self.profile,dict(id=OTHER,display_name='Researcher B')],parent)
        editor.attendance.cellWidget(0,1).setChecked(True)
        editor.notes.setPlainText('Decisions persisted')
        editor.transcript.setPlainText(import_transcript(transcript))
        editor.add_action(dict(id=str(uuid.uuid4()),title='Check pore map',assigned_to=USER,due_date='2026-10-16',completed=True))
        editor.files.setCurrentRow(0)
        with patch('app.ui.participation.file_launcher.open_file') as opened,patch('app.ui.participation.file_launcher.show_in_folder') as shown:
            editor.open_file(False);editor.open_file(True);opened.assert_called_once_with(audio);shown.assert_called_once_with(audio)
        editor.save()
        loaded=self.service.get(row['id']);self.assertTrue(loaded['attendees'][0]['attended']);self.assertFalse(loaded['attendees'][1]['attended'])
        self.assertTrue(loaded['actions'][0]['completed']);self.assertEqual(loaded['content']['notes'],'Decisions persisted')
        self.assertIn('Holy GDL',self.service.list('Holy')[0]['content']['transcript'])
        editor.deleteLater();parent.deleteLater();QApplication.processEvents()

    def test_pdf_csv_export_person_week_and_no_credentials(self):
        from pypdf import PdfReader
        begin=datetime(2026,10,5,tzinfo=timezone.utc)
        reports=[aggregate(dict(id=identity,display_name=name),begin,begin+timedelta(days=7),[],[],[],[])
                 for identity,name in [(USER,'Researcher A'),(OTHER,'Researcher B')]]
        for report in reports:
            pdf=self.fixture.root/(report['person']+'.pdf');csv=self.fixture.root/(report['person']+'.csv')
            export_pdf(report,pdf);export_csv(report,csv)
            text='\n'.join(p.extract_text() for p in PdfReader(pdf).pages)
            self.assertIn(report['person'],text);self.assertIn('2026-10-05',text)
            self.assertNotIn('access_token',text);self.assertIn('Files opened',csv.read_text(encoding='utf-8-sig'))
            self.assertNotIn('Researcher B' if report['person']=='Researcher A' else 'Researcher A',text)

    def test_generate_all_team_ui_and_self_reporting_permissions(self):
        from app.ui.participation import WeeklyPanel
        import time
        normal=WeeklyReports(self.accounts);begin=datetime.now().astimezone()-timedelta(days=1);end=begin+timedelta(days=7)
        with self.assertRaises(ValueError):normal.load(dict(id=OTHER),begin,end)
        self.accounts.profile['role']='admin'
        parent=QMainWindow();panel=WeeklyPanel(parent,self.accounts)
        def settle():
            limit=time.monotonic()+5
            while panel.tasks.busy() and time.monotonic()<limit:QApplication.processEvents();time.sleep(.01)
            self.assertFalse(panel.tasks.busy())
        QApplication.processEvents();settle()
        with patch('app.ui.participation.QFileDialog.getExistingDirectory',return_value=str(self.fixture.root)):
            panel.export_all();settle()
        self.assertEqual(len(panel.reports),2)
        files=list(self.fixture.root.glob('Researcher*.pdf'));self.assertEqual(len(files),2)
        self.assertEqual({r['person'] for r in panel.reports},{'Researcher A','Researcher B'})
        self.assertIn('2 individual team reports exported',panel.status.text())
        panel.deleteLater();parent.deleteLater();QApplication.processEvents()
