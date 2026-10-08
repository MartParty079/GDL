"""Research object, file association, migration, OS action and viewer contracts."""

import json
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from app.services.storage import Store, ROOT
from app.services.project_locations import ProjectLocations
from app.indexing.index_manager import NativeIndex
from app.services.research_workspace import ResearchWorkspace
from app.services import file_launcher, previews


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.current = self.root / "current"
        self.legacy = self.root / "legacy"
        self.current.mkdir()
        self.legacy.mkdir()
        self.store = Store(ROOT, self.root / "profile")
        self.locations = ProjectLocations(self.store)
        value = self.locations.value
        value["active"]["root_path"] = str(self.current)
        value["legacy"][0]["root_path"] = str(self.legacy)
        value["enabled"] = True
        self.locations.save(value)
        self.catalog = NativeIndex(self.store, self.locations)
        self.repo = ResearchWorkspace(self.catalog)
        (self.current / "GDL-003.txt").write_text("research fixture", encoding="utf-8")
        (self.legacy / "GDL-003.txt").write_text("historical fixture", encoding="utf-8")
        self.catalog.refresh()
        self.row = self.catalog.query(origin="current")[0][0]
        self.old = self.catalog.query(origin="legacy")[0][0]

    def tearDown(self):
        self.temp.cleanup()

    def sample(self):
        return self.repo.save(
            "sample",
            {"name": "Carbon paper", "material": "carbon", "notes": "microscopy notes"},
            "GDL-003",
        )

    def test_flexible_objects_associations_search_timeline_and_revisions(self):
        self.sample()
        exp = self.repo.save(
            "experiment",
            {
                "name": "Compression cycle",
                "sample_id": "GDL-003",
                "conditions": [{"name": "Pressure", "value": "20", "unit": "kPa"}],
                "results": [{"name": "Porosity", "value": "71", "unit": "%"}],
            },
            "EXP-001",
        )
        self.assertEqual(exp["results"][0]["unit"], "%")
        self.assertEqual(exp["date"], "")
        self.repo.annotate(
            [self.row],
            {"experiment_id": "EXP-001", "tags": "fixture", "report_version": "v2"},
        )
        rows, total = self.catalog.query(
            "compression", sample="GDL-003", experiment="EXP-001", origin="current"
        )
        self.assertEqual(total, 1)
        self.assertEqual(rows[0]["id"], self.row["id"])
        self.assertEqual(len(self.repo.revisions(self.row)), 1)
        self.assertTrue(self.repo.timeline(sample="GDL-003"))
        self.assertTrue(self.repo.timeline(experiment="EXP-001"))
        self.assertEqual((self.current / "GDL-003.txt").read_text(), "research fixture")
        self.assertEqual(self.repo.get("GDL-003")["received_date"], "")
        self.assertEqual(self.catalog.query(empty_context=True)[1], 0)
        self.repo.save(
            "sample",
            {"name": "Carbon paper", "description": "new searchable words"},
            "GDL-003",
        )
        self.assertEqual(self.catalog.query("searchable", sample="GDL-003")[1], 1)

    def test_cross_origin_and_inconsistent_sample_bulk_roll_back(self):
        self.sample()
        with self.assertRaises(ValueError):
            self.repo.annotate([self.row, self.old], {"sample_id": "GDL-003"})
        self.assertEqual(self.catalog.query(sample="GDL-003")[1], 0)
        self.repo.save("sample", {"name": "Other"}, "GDL-004")
        self.repo.save("experiment", {"name": "Run", "sample_id": "GDL-003"}, "EXP-001")
        with self.assertRaises(ValueError):
            self.repo.annotate(
                [self.row], {"sample_id": "GDL-004", "experiment_id": "EXP-001"}
            )
        self.assertEqual(self.repo.revisions(self.row), [])

    def test_suggestions_do_not_assign_and_legacy_objects_remain_distinct(self):
        self.sample()
        self.repo.save(
            "sample", {"name": "Carbon paper"}, "LEGACY-GDL-003", origin="legacy"
        )
        self.assertEqual(self.repo.suggestions(self.row)[0]["id"], "GDL-003")
        self.assertEqual(self.catalog.query(sample="GDL-003")[1], 0)
        self.assertEqual(len(self.repo.objects("sample", "legacy")), 1)
        with self.assertRaises(ValueError):
            self.repo.save("sample", {"name": "Duplicate"}, "GDL-003", origin="legacy")

    def test_global_sql_sort_combined_filters_and_pagination(self):
        for n in range(205):
            (self.current / f"file-{n:03}.csv").write_text("col\n1")
        self.catalog.refresh()
        rows, total = self.catalog.query(
            origin="current", family="Data", sort="Name", offset=200, limit=200
        )
        self.assertEqual(total, 205)
        self.assertEqual(rows[0]["name"], "file-200.csv")
        self.assertEqual(len(rows), 5)
        self.repo.viewed(self.row)
        self.repo.annotate([self.row], {"favorite": True})
        self.assertEqual(self.catalog.query(recent=True)[1], 1)
        self.assertEqual(self.catalog.query(favorite=True)[1], 1)
        for sort in ("Created newest", "Recently added", "Report version", "Status"):
            self.assertEqual(self.catalog.query(sort=sort)[1], 207)

    def test_metadata_tags_append_and_index_refresh_preserve_objects(self):
        self.sample()
        self.repo.annotate([self.row], {"sample_id": "GDL-003", "tags": "one"})
        self.repo.annotate([self.row], {"tags": "two, one"}, append_tags=True)
        self.catalog.refresh()
        row = self.catalog.query(sample="GDL-003")[0][0]
        self.assertEqual(row["tags"], "one, two")
        self.assertEqual(self.repo.get("GDL-003")["name"], "Carbon paper")
        self.assertFalse(
            any(e["kind"] == "Filesystem file modified" for e in self.repo.timeline())
        )

    def test_v3_migration_preserves_ids_content_overrides_and_makes_backup(self):
        self.catalog.set_override(self.row, {"notes": "manual protected"})
        with self.catalog.connect() as db:
            before = db.execute("SELECT id,payload FROM files ORDER BY id").fetchall()
            fts = db.execute(
                "SELECT rowid,* FROM content_fts ORDER BY rowid"
            ).fetchall()
            db.execute("PRAGMA user_version=3")
        with patch.object(
            NativeIndex,
            "upsert",
            side_effect=AssertionError("v3 migration must not rebuild file records"),
        ):
            NativeIndex(self.store, self.locations)
        with self.catalog.connect() as db:
            self.assertEqual(
                db.execute("SELECT id,payload FROM files ORDER BY id").fetchall(),
                before,
            )
            self.assertEqual(
                db.execute("SELECT rowid,* FROM content_fts ORDER BY rowid").fetchall(),
                fts,
            )
            self.assertEqual(db.execute("PRAGMA user_version").fetchone()[0], 4)
        self.assertTrue(
            list(
                Path(self.locations.value["backups"]).glob("before-migration-*.sqlite3")
            )
        )

    def test_locate_missing_record_safe_and_hide_is_reversible(self):
        path = self.current / "renamed.txt"
        (self.current / "GDL-003.txt").rename(path)
        updated = self.repo.locate(self.row, path)
        self.assertEqual(updated["id"], self.row["id"])
        self.assertEqual(updated["name"], "renamed.txt")
        self.repo.hide(updated)
        self.assertEqual(self.catalog.query(origin="current")[1], 0)
        self.repo.restore_hidden()
        self.assertEqual(self.catalog.query(origin="current")[1], 1)
        with self.assertRaises(ValueError):
            self.repo.locate(updated, self.legacy / "GDL-003.txt")

    def test_research_events_dates_and_object_kinds_are_validated(self):
        self.sample()
        self.repo.record_event("Note added", "Prepared for imaging", sample="GDL-003")
        self.assertEqual(self.repo.timeline(sample="GDL-003")[0]["kind"], "Note added")
        with self.assertRaises(ValueError):
            self.repo.save(
                "sample", {"name": "Bad date", "received_date": "yesterday"}, "S-BAD"
            )
        self.repo.save("experiment", {"name": "Run", "sample_id": "GDL-003"}, "EXP-001")
        with self.assertRaises(ValueError):
            self.repo.annotate([self.row], {"sample_id": "EXP-001"})
        self.repo.annotate([self.row], {"experiment_id": "EXP-001"})
        self.repo.save("sample", {"name": "Second sample"}, "GDL-004")
        with self.assertRaises(ValueError):
            self.repo.save(
                "experiment", {"name": "Run", "sample_id": "GDL-004"}, "EXP-001"
            )

    def test_filename_inference_never_becomes_a_confirmed_link(self):
        with self.catalog.connect() as db:
            self.catalog.upsert(
                db, {**self.row, "sample_id": "GDL-003"}, update_search=True
            )
        self.sample()
        self.assertEqual(self.catalog.query(sample="GDL-003", explicit=True)[1], 0)
        self.assertEqual(self.catalog.query("Carbon paper")[1], 0)
        self.repo.annotate([self.row], {"sample_id": "GDL-003"})
        self.assertEqual(self.catalog.query(sample="GDL-003", explicit=True)[1], 1)
        self.assertEqual(self.catalog.query("Carbon paper")[1], 1)
        with self.catalog.connect() as db:
            self.catalog.upsert(
                db, {**self.old, "sample_id": "GDL-003"}, update_search=True
            )
        self.repo.annotate([self.old], {"tags": "Legacy notes retained"})
        self.assertEqual(self.catalog.query(origin="legacy", tags="Legacy notes")[1], 1)

    def test_windows_associations_and_explorer_preserve_full_arguments(self):
        for extension in (".xlsx", ".pdf", ".docx", ".png", ".csv"):
            path = self.current / ("space & quote " + extension)
            path.write_bytes(b"fixture")
            with patch.object(file_launcher.sys, "platform", "win32"), patch.object(
                file_launcher.os, "startfile", create=True
            ) as launch:
                file_launcher.open_file(path)
                launch.assert_called_once_with(str(path))
        with patch.object(file_launcher.sys, "platform", "win32"), patch.object(
            file_launcher.subprocess, "Popen"
        ) as launch:
            file_launcher.show_in_folder(self.current / "GDL-003.txt")
            launch.assert_called_once_with(
                ["explorer.exe", "/select,", str(self.current / "GDL-003.txt")]
            )
        with patch.object(file_launcher.sys, "platform", "win32"), patch.object(
            file_launcher.os, "startfile", create=True
        ) as launch:
            file_launcher.open_folder(self.current)
            launch.assert_called_once_with(str(self.current))
        with self.assertRaises(ValueError):
            file_launcher.open_file(self.current / "missing.csv")

    def test_python_never_executes_and_escape_path_rejected(self):
        path = self.current / "code.py"
        path.write_text('raise Exception("never run")')
        with patch.object(file_launcher.sys, "platform", "win32"), patch.object(
            file_launcher.subprocess, "Popen"
        ) as launch:
            file_launcher.open_file(path)
            launch.assert_called_once_with(["notepad.exe", str(path)])
        with self.assertRaises(ValueError):
            self.catalog.safe_path({**self.row, "relative_path": "../outside.txt"})

    def test_bounded_resident_previews_and_cache_do_not_modify_sources(self):
        from PIL import Image
        from openpyxl import Workbook
        from docx import Document

        path = self.current / "image.png"
        Image.new("RGB", (80, 60), "red").save(path)
        book = Workbook()
        book.active.append(["value", "=1+1"])
        for n in range(500):
            book.active.append([n, n])
        book.save(self.current / "data.xlsx")
        book.close()
        doc = Document()
        doc.add_paragraph("Readable report")
        doc.save(self.current / "report.docx")
        self.catalog.refresh()
        rows = {r["name"]: r for r in self.catalog.query(origin="current")[0]}
        before = path.read_bytes()
        cache = Path(previews.thumbnail(self.catalog, rows["image.png"]))
        self.assertTrue(cache.is_relative_to(Path(self.locations.value["cache"])))
        self.assertEqual(path.read_bytes(), before)
        data = previews.table(self.catalog, rows["data.xlsx"])
        self.assertEqual(len(data), 250)
        self.assertEqual(data[0][1], "=1+1")
        self.assertIn(
            "Readable report", previews.text(self.catalog, rows["report.docx"])
        )
        with patch(
            "app.services.previews.placeholder", return_value=True
        ), patch.object(Path, "open", side_effect=AssertionError("no hydration")):
            with self.assertRaisesRegex(ValueError, "Online-only"):
                previews.thumbnail(self.catalog, rows["image.png"])


class WorkspaceUITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def wait(self, hub):
        deadline = time.monotonic() + 20
        while hub.busy():
            self.app.processEvents()
            time.sleep(0.005)
            if time.monotonic() > deadline:
                self.fail("Workspace background tasks timed out")
        self.app.processEvents()

    def test_sample_experiment_scoped_tabs_preview_and_version(self):
        from app.ui.window import HubWindow
        from app import __version__
        from PySide6.QtCore import QCoreApplication, QEvent

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            current = root / "current"
            legacy = root / "legacy"
            current.mkdir()
            legacy.mkdir()
            store = Store(ROOT, root / "profile")
            locations = ProjectLocations(store)
            value = locations.value
            value["active"]["root_path"] = str(current)
            value["legacy"][0]["root_path"] = str(legacy)
            value["enabled"] = True
            locations.save(value)
            (current / "GDL-003.csv").write_text("pressure,value\n1,2")
            from pypdf import PdfWriter
            from PIL import Image

            pdf = PdfWriter()
            pdf.add_blank_page(width=200, height=300)
            pdf.write(current / "viewer.pdf")
            Image.new("RGB", (80, 60), "red").save(current / "viewer.png")

            index = NativeIndex(store, locations)
            index.refresh()
            repo = ResearchWorkspace(index)
            repo.save("sample", {"name": "Test carbon"}, "GDL-003")
            repo.save(
                "experiment",
                {"name": "Pressure test", "sample_id": "GDL-003"},
                "EXP-003",
            )
            row = index.query()[0][0]
            repo.annotate([row], {"experiment_id": "EXP-003"})
            store.local.update(research_workspace_sizes=[0,1280,0],
                               research_navigation_visible=True,research_details_visible=True)
            with patch("app.ui.window.QTimer.singleShot"):
                window = HubWindow(store)
            hub = window.research_workspace
            window.resize(1280, 850)
            window.show()
            self.app.processEvents()
            try:
                from PySide6.QtWidgets import QPushButton
                self.assertIn("v"+__version__, [control.text() for control in window.findChildren(QPushButton)])
                self.assertGreater(hub.splitter.sizes()[0],0)
                self.assertGreaterEqual(hub.splitter.sizes()[2],260)
                self.assertEqual(window.project_tabs.currentIndex(), 1)
                self.assertEqual(window.tabs.currentIndex(), 3)
                hub.navigate("Samples")
                self.assertEqual(hub.object_list.count(), 1)
                hub.show_object("GDL-003")
                self.assertEqual(hub.object_tabs.count(), 7)
                hub.object_tabs.setCurrentIndex(3)
                self.wait(hub)
                browser = hub.object_tabs.currentWidget()
                self.assertEqual(browser.total, 1)
                hub.show_object("EXP-003")
                self.assertEqual(hub.object_tabs.tabText(4), "Results")
                hub.object_tabs.setCurrentIndex(1)
                self.wait(hub)
                self.assertEqual(hub.object_tabs.currentWidget().total, 1)
                hub.select_file(index.query()[0][0])
                self.wait(hub)
                self.assertEqual(hub.preview.table.rowCount(), 2)
                from PySide6.QtPdf import QPdfDocument

                pdfrow = next(r for r in index.query()[0] if r["name"] == "viewer.pdf")
                imagerow = next(
                    r for r in index.query()[0] if r["name"] == "viewer.png"
                )
                hub.select_file(pdfrow)
                self.wait(hub)
                self.assertEqual(len(hub.preview.content.findChildren(QPdfDocument)), 1)
                hub.select_file(row)
                self.wait(hub)
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                self.assertEqual(hub.preview.content.findChildren(QPdfDocument), [])
                hub.file_action("preview", imagerow)
                self.wait(hub)
                viewer = hub.dialogs[-1]
                viewer.reject()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                self.assertEqual(hub.dialogs, [])
                hub.navigate("Files")
                self.wait(hub)
                self.assertGreater(hub.browser.table.height(), 180)
                hub.browser.table.selectRow(0)
                hub.file_action("copy", row)
                from PySide6.QtWidgets import QApplication

                self.assertEqual(
                    QApplication.clipboard().text(), str(current / "GDL-003.csv")
                )
            finally:
                self.wait(hub)
                window.close()
                window.deleteLater()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
                self.app.processEvents()
