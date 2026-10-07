"""Local indexing contracts: extraction, safe search, incremental work and offline UI."""
import builtins
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from app.services.storage import Store, ROOT
from app.services.project_locations import ProjectLocations
from app.indexing.index_manager import NativeIndex
from app.indexing.extractor import extract
from app.services.project_storage import IndexCancelled


class NativeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.current = self.root / 'current'
        self.legacy = self.root / 'legacy'
        self.current.mkdir(); self.legacy.mkdir()
        self.store = Store(ROOT, self.root / 'profile')
        self.locations = ProjectLocations(self.store)
        value = self.locations.value
        value['active']['root_path'] = str(self.current)
        value['legacy'][0]['root_path'] = str(self.legacy)
        self.locations.save(value)
        self.index = NativeIndex(self.store, self.locations)

    def tearDown(self):
        self.temp.cleanup()

    def test_text_formats_and_binary_rejection_are_bounded(self):
        for suffix in ('.txt', '.md', '.csv', '.json', '.py', '.m', '.xml', '.yaml', '.yml'):
            path = self.current / ('fixture' + suffix)
            path.write_text('pore diameter porosity ' * 100)
            result = extract(path, {'max_text_chars': 25})
            self.assertEqual(len(result['text']), 25)
            self.assertEqual(result['status'], 'Partial')
        path = self.current / 'binary.txt'
        path.write_bytes(b'abc\x00xyz')
        self.assertEqual(extract(path)['status'], 'Unsupported binary content')
        self.assertEqual(extract(path, {'enabled': False})['status'], 'Extraction disabled')
        path.write_bytes('microscopy'.encode('utf-16'))
        self.assertIn('microscopy', extract(path)['text'])

    def test_docx_xlsx_pptx_and_header_image(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from PIL import Image
        path = self.current / 'report.docx'
        doc = Document(); doc.core_properties.title = 'Microscopy'; doc.add_paragraph('fiber diameter')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'roundness'
        doc.save(path)
        result = extract(path)
        self.assertIn('roundness', result['text']); self.assertEqual(result['metadata']['title'], 'Microscopy')
        path = self.current / 'workbook.xlsx'
        book = Workbook(); sheet = book.active; sheet.title = 'Water Intrusion'
        sheet.append(['pressure', 'syringe', '=SECRET_FORMULA()', 999.25]); book.save(path); book.close()
        result = extract(path)
        self.assertIn('syringe', result['text']); self.assertNotIn('SECRET_FORMULA', result['text'])
        self.assertNotIn('999.25', result['text']); self.assertEqual(result['metadata']['sheets']['Water Intrusion']['columns'], 4)
        self.assertEqual(extract(path, {'max_cells': 1})['status'], 'Partial')
        path = self.current / 'slides.pptx'
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[0])
        slide.shapes.title.text = 'ASTM compression'; slides.save(path)
        result = extract(path)
        self.assertIn('ASTM', result['text']); self.assertEqual(result['metadata']['slide_count'], 1)
        path = self.current / 'image.png'; Image.new('RGB', (17, 29)).save(path)
        with patch.object(Image.Image, 'load', side_effect=AssertionError('pixel decoding')):
            result = extract(path)
        self.assertEqual(result['metadata']['width'], 17); self.assertEqual(result['metadata']['height'], 29)

    def test_pdf_text_metadata_and_ocr_required(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
        writer.add_metadata({'/Title': 'Research pore report'})
        path = self.current / 'blank.pdf'; writer.write(path)
        self.assertEqual(extract(path)['status'], 'OCR required')
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 20 Td (equivalent pore diameter) Tj ET')
        page[NameObject('/Contents')] = stream; writer.write(path)
        result = extract(path)
        self.assertIn('equivalent pore diameter', result['text'])
        self.assertEqual(result['metadata']['page_count'], 1); self.assertEqual(result['metadata']['title'], 'Research pore report')

    def test_corrupt_and_unsupported_do_not_abort_scan(self):
        (self.current / 'bad.pdf').write_bytes(b'not a PDF')
        (self.current / 'shape.dwg').write_bytes(b'proprietary')
        (self.current / 'good.md').write_text('solidity porosity')
        result = self.index.refresh()
        self.assertEqual(result['scanned'], 3); self.assertEqual(result['errors'], 1)
        self.assertEqual(self.index.query(query='solidity')[1], 1)
        self.assertEqual(self.index.history()[0]['result'], 'Completed with Errors')
        with self.index.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM index_errors').fetchone()[0], 1)

    def test_safe_ranked_fts_prefix_metadata_and_filters(self):
        (self.current / 'porosity.txt').write_text('pore diameter')
        (self.legacy / 'other.txt').write_text('porosity roundness')
        (self.current / 'other.txt').write_text('porosity roundness')
        self.index.refresh()
        rows, count = self.index.query(query='poros', origin='current')
        self.assertEqual(count, 2); self.assertEqual(rows[0]['name'], 'porosity.txt')
        self.assertEqual(self.index.query(query='round', origin='legacy')[1], 1)
        self.assertEqual(self.index.query(query='" OR * NOT (')[1], 0)
        item = rows[0]
        self.index.set_override(item, {'tags': 'syringe', 'description': 'laser fixture', 'favorite': True, 'category': 'Standards'})
        self.assertEqual(self.index.query(query='laser', favorite=True, category='Standards')[1], 1)
        self.assertEqual(self.index.query(tags='syringe', extension='.txt')[1], 1)
        self.index.refresh(rebuild=True)
        self.assertEqual(self.index.query(query='laser', favorite=True)[1], 1)

    def test_no_content_reads_for_unchanged_or_search_rebuild(self):
        (self.current / 'report.txt').write_text('fiber diameter')
        self.index.refresh()
        with patch('app.indexing.index_manager.extract', side_effect=AssertionError('unchanged file reopened')):
            self.assertEqual(self.index.refresh(quick=True)['unchanged'], 1)
            self.index.rebuild_search()
        self.assertEqual(self.index.query(query='fiber')[1], 1)
        self.assertEqual(self.index.history()[0]['mode'], 'search')
        self.store.local['native_indexing'] = {'limits': {'max_text_chars': 4}}
        self.index.refresh(quick=True)
        self.assertEqual(self.index.query(query='diameter')[1], 0)

    def test_online_only_does_not_open_a_parser_or_hash(self):
        path = self.current / 'cloud.docx'; path.write_bytes(b'placeholder')
        with patch('app.indexing.index_manager.placeholder', return_value=True), patch('app.indexing.index_manager.extract', side_effect=AssertionError('hydration')):
            self.index.refresh()
        self.assertEqual(self.index.summary()['online_only'], 1)
        with patch('app.indexing.extractor.placeholder', return_value=True), patch.object(Path, 'open', side_effect=AssertionError('hydration')):
            self.assertEqual(extract(path)['status'], 'Online-only')

    def test_sql_pagination_relationships_backup_restore_and_cancel(self):
        for n in range(260):
            (self.current / f'item{n:03}.txt').write_text('pore')
        self.index.refresh()
        first, total = self.index.query(limit=200); second, _ = self.index.query(offset=200)
        self.assertEqual(total, 260); self.assertEqual(len(second), 60)
        self.assertFalse({r['id'] for r in first} & {r['id'] for r in second})
        self.index.relate(first[0]['id'], second[0]['id'], 'derived_from')
        self.assertEqual(len(self.index.related(first[0]['id'])), 1)
        backup = self.index.backup()
        self.index.set_override(first[0], {'tags': 'new'})
        self.index.restore(backup)
        self.assertEqual(self.index.query(tags='new')[1], 0)
        with self.assertRaises(IndexCancelled):
            self.index.refresh(cancel=lambda: True)
        self.assertEqual(self.index.history()[0]['result'], 'Cancelled')
        self.assertEqual(self.index.query()[1], 260)

    def test_additional_sources_keep_current_reference_and_archive_separate(self):
        import copy
        import uuid
        value = copy.deepcopy(self.locations.value)
        value['additional'] = []
        for kind in ('active', 'reference', 'archive'):
            folder = self.root / kind; folder.mkdir(); (folder / 'record.txt').write_text(kind)
            source = dict(value['active'], id=str(uuid.uuid4()), name=kind, source_label=kind,
                root_path=str(folder), source_type=kind, project_type=kind,
                read_only=kind != 'active', dataset_status=kind)
            value['additional'].append(source)
        self.locations.save(value); self.index.sync_sources(); self.index.refresh()
        self.assertEqual(self.index.query(origin='current')[1], 1)
        self.assertEqual(self.index.query(origin='reference')[1], 1)
        self.assertEqual(self.index.query(origin='archive', archive='archived')[1], 1)
        self.assertEqual(self.index.query(archive='all')[1], 3)

    def test_cancel_retains_committed_batches_without_marking_missing(self):
        for n in range(400):
            (self.current / f'{n:03}.txt').write_text('fixture')
        calls = 0
        def cancel():
            nonlocal calls
            calls += 1
            return calls >= 300
        with self.assertRaises(IndexCancelled):
            self.index.refresh(cancel=cancel)
        self.assertEqual(self.index.query()[1], 250)
        self.assertTrue(all(row['availability'] == 'Locally available' for row in self.index.query()[0]))
        self.index.refresh(quick=True)
        self.assertEqual(self.index.query()[1], 400)

    def test_unavailable_log_folder_and_empty_source_selection_are_safe(self):
        cache = self.root / 'cache-conflict'; cache.write_text('unavailable directory')
        self.locations.value['cache'] = str(cache)
        (self.current / 'record.txt').write_text('data')
        self.assertEqual(self.index.refresh(source_ids=[])['scanned'], 0)
        self.assertEqual(self.index.query()[1], 0)
        self.assertEqual(self.index.refresh()['scanned'], 1)
        self.assertTrue(self.index.logging_unavailable)
        self.assertEqual(self.index.history()[0]['result'], 'Complete')

    def test_native_offline_startup_without_auth_packages(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QCoreApplication, QEvent
        from app.ui.window import HubWindow
        app = QApplication.instance() or QApplication([])
        original = builtins.__import__
        def guarded(name, *args, **kwargs):
            if name.startswith(('msal', 'azure.identity', 'app.services.microsoft')):
                raise AssertionError('retired authentication import')
            return original(name, *args, **kwargs)
        with patch('builtins.__import__', guarded), patch('socket.socket', side_effect=AssertionError('network')), patch('app.ui.window.QTimer.singleShot'):
            window = HubWindow(self.store)
            self.assertFalse(hasattr(window, 'auth'))
            names = [window.settings_tabs.tabText(n) for n in range(window.settings_tabs.count())]
            self.assertFalse(any('Microsoft' in name or 'Cloud' in name for name in names))
            window.close(); window.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


class WatcherTests(unittest.TestCase):
    def test_burst_debounce_periodic_fallback_and_disable(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtWidgets import QApplication
        from PySide6.QtTest import QTest
        from app.indexing.watcher import IndexWatcher
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as temporary:
            watcher = IndexWatcher()
            watcher.configure([temporary], watching=True, automatic=False, minutes=2)
            self.assertTrue(watcher.periodic.isActive())
            self.assertIn(temporary, watcher.watcher.directories())
            calls = []
            watcher.refresh_requested.connect(lambda: calls.append(True))
            watcher.debounce.setInterval(20)
            watcher.schedule(); watcher.schedule(); watcher.schedule()
            QTest.qWait(80)
            self.assertEqual(len(calls), 1)
            watcher.schedule()
            watcher.configure([], watching=False, automatic=False)
            QTest.qWait(80)
            self.assertEqual(len(calls), 1)
            self.assertFalse(watcher.periodic.isActive())
            watcher.stop()
