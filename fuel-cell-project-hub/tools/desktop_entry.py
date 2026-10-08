"""Windows package entry and isolated startup acceptance check."""
import sys
import tempfile
from pathlib import Path

if not getattr(sys, 'frozen', False):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def self_check(report):
    import json
    from unittest.mock import patch
    from PySide6.QtCore import QCoreApplication, QEvent
    from PySide6.QtWidgets import QApplication
    from app.services.storage import Store, ROOT, write_json
    from app.ui.window import HubWindow
    from app.ui.branding import application_icon
    from app.services.resources import resource_path
    with tempfile.TemporaryDirectory() as temporary:
        app = QApplication([])
        app.setApplicationName('FuelCellProjectHub')
        app.setWindowIcon(application_icon())
        with patch('app.ui.window.QTimer.singleShot'):
            window = HubWindow(Store(ROOT, Path(temporary) / 'profile'))
        window.show()
        app.processEvents()
        result = {'startup': True, 'settings_tabs': [window.settings_tabs.tabText(i) for i in range(window.settings_tabs.count())],
                  'application_icon_loaded': not app.windowIcon().isNull(), 'window_icon_loaded': not window.windowIcon().isNull(),
                  'icon_asset_present': resource_path('assets/app_icon.ico').is_file(),
                  'local_path_editable': not window.storage_panel.root.isReadOnly(), 'default_origin': window.files_panel.origin.currentData(),
                  'engine_manifest_present': (ROOT / 'analysis/gdl/manifest.json').is_file(), 'fiji_launched': False}
        current, legacy = Path(temporary) / 'current', Path(temporary) / 'legacy'
        current.mkdir()
        legacy.mkdir()
        (current / 'current.txt').write_text('current porosity')
        (legacy / 'historical.txt').write_text('legacy')
        panel = window.research_panel
        panel.active.setText(str(current))
        panel.legacy.setPlainText(str(legacy))
        panel.paths['shared_storage'].setText(str(current))
        if not panel.save():
            raise ValueError('Packaged research configuration failed.')
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pypdf import PdfWriter
        from PIL import Image
        document = Document(); document.add_paragraph('fiber diameter'); document.save(current / 'fixture.docx')
        workbook = Workbook(); workbook.active.append(['water intrusion']); workbook.save(current / 'fixture.xlsx'); workbook.close()
        presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[0]); slide.shapes.title.text = 'compression'; presentation.save(current / 'fixture.pptx')
        pdf = PdfWriter(); pdf.add_blank_page(width=100, height=100); pdf.write(current / 'fixture.pdf')
        Image.new('RGB', (10, 20)).save(current / 'fixture.png')
        indexed = window.storage_settings.catalog.refresh()
        catalog = window.storage_settings.catalog
        assert indexed['errors'] == 0
        assert catalog.query(query='fiber diameter')[1] == 1
        assert catalog.query(query='water intrusion')[1] == 1
        assert catalog.query(query='compression')[1] == 1
        assert not hasattr(window, 'auth')
        result.update(native_content_search=True, local_parsers_present=True, graph_storage_removed=True,
                      microsoft_identity_supported=True)
        window.files_panel.reload()
        result.update(research_catalog_available=True, research_current_files=indexed['current'],
                      research_legacy_files=indexed['legacy'], research_index_errors=indexed['errors'],
                      research_template_present=resource_path('config/research_defaults.json').is_file(),
                      research_sources_unchanged=not (current / '.projecthub').exists() and not (legacy / '.projecthub').exists())
        from app import __version__
        from app.services.research_workspace import ResearchWorkspace
        from app.services import previews
        import time
        repository=ResearchWorkspace(catalog)
        repository.save('sample', {'name':'Packaged sample'}, 'GDL-003')
        repository.save('experiment', {'name':'Packaged experiment','sample_id':'GDL-003','conditions':[{'name':'Pressure','value':'20','unit':'kPa'}]}, 'EXP-003')
        row=catalog.query(origin='current')[0][0]
        repository.annotate([row], {'experiment_id':'EXP-003'})
        assert catalog.query(sample='GDL-003',experiment='EXP-003',explicit=True)[1]==1
        workspace=window.research_workspace
        workspace.show_object('GDL-003')
        assert workspace.object_tabs.count()==7
        workspace.show_object('EXP-003')
        assert workspace.object_tabs.tabText(4)=='Results'
        rows={r['name']:r for r in catalog.query(origin='current')[0]}
        assert len(previews.table(catalog,rows['fixture.xlsx']))==1
        assert Path(previews.thumbnail(catalog,rows['fixture.png'])).is_file()
        from PySide6.QtPdf import QPdfDocument
        document=QPdfDocument(window)
        assert document.load(str(previews.resident_path(catalog,rows['fixture.pdf'])))==QPdfDocument.Error.None_
        assert document.pageCount()==1
        assert resource_path('CHANGELOG.md').is_file()
        from app.services.storage import read_json
        assert read_json(resource_path('version_history.json'),{})['versions'][-1]['version']==__version__
        if '--local-only' not in sys.argv:
            from app.services.desktop_oauth import callback_values, CALLBACK, CallbackBroker
            assert callback_values(CALLBACK+'?code=packaged-check')['code']=='packaged-check'
            broker=CallbackBroker(Path(temporary)/'callback-profile')
            assert broker.listen()
            broker.server.close()
            result.update(protocol_callback_available=True)
        result.update(structured_version_history_available=True)
        result.update(version=__version__, research_objects_available=True, rendered_pdf_available=True,
                      spreadsheet_preview_available=True, image_thumbnail_available=True, changelog_present=True)
        from app.services.accounts import Accounts, protect
        from app.ui.accounts import LoginDialog
        import os
        account_store=Store(ROOT,Path(temporary)/'account-profile')
        accounts=Accounts(account_store,lambda *args: None)
        login=LoginDialog(accounts)
        deadline=time.monotonic()+10
        while login.tasks.busy():
            app.processEvents(); time.sleep(.01)
            if time.monotonic()>deadline: raise RuntimeError('Login startup timed out')
        app.processEvents()
        while login.tasks.busy(): app.processEvents();time.sleep(.01)
        assert login.signin.isEnabled()
        assert account_store.config_dir.joinpath('accounts_public.json').is_file()
        if os.name=='nt': assert protect(protect(b'isolated-session'),True)==b'isolated-session'
        login.close(); accounts.executor.shutdown()
        result.update(login_ui_available=True, client_configuration_bundled=True,
                      windows_session_encryption_available=os.name=='nt')
        deadline=time.monotonic()+30
        while workspace.busy():
            app.processEvents();time.sleep(.01)
            if time.monotonic()>deadline:raise ValueError('Packaged workspace tasks did not finish.')
        app.processEvents()
        window.close()
        window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        write_json(Path(report), result)
    return 0


if __name__ == '__main__':
    if '--self-check' in sys.argv:
        report_path = sys.argv[sys.argv.index('--self-check') + 1]
        try:
            sys.exit(self_check(report_path))
        except Exception as exc:
            import json
            Path(report_path).write_text(json.dumps({'startup': False, 'error': str(exc)}), encoding='utf-8')
            sys.exit(1)
    from app.main import main
    sys.exit(main())
