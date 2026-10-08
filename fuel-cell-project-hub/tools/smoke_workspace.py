"""Offscreen research acceptance with isolated fixtures, or read-only source browsing."""

import os
import sys
import tempfile
import time
import json
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["QT_QPA_PLATFORM"] = "offscreen"
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import QCoreApplication, QEvent
from app.services.storage import Store, ROOT
from app.services.project_locations import ProjectLocations
from app.indexing.index_manager import NativeIndex
from app.services.research_workspace import ResearchWorkspace
from app.ui.window import HubWindow


def wait(app, hub):
    deadline = time.monotonic() + 30
    idle_since = None
    while True:
        app.processEvents()
        if hub.busy():
            idle_since = None
        elif idle_since is None:
            idle_since = time.monotonic()
        elif time.monotonic() - idle_since > 0.1:
            break
        time.sleep(0.01)
        if time.monotonic() > deadline:
            raise RuntimeError("Workspace tasks timed out")


def run():
    app = QApplication([])
    temporary = tempfile.TemporaryDirectory()
    if "--personal" in sys.argv:
        store = Store()
    else:
        from PIL import Image, ImageDraw
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter

        root = Path(temporary.name)
        current = root / "current"
        legacy = root / "legacy"
        current.mkdir()
        legacy.mkdir()
        store = Store(ROOT, root / "profile")
        locations = ProjectLocations(store)
        value = locations.value
        value["active"]["root_path"] = str(current)
        value["legacy"][0]["root_path"] = str(legacy)
        value["enabled"] = True
        locations.save(value)
        for number in range(6):
            image = Image.new("RGB", (700, 500), "#222f3a")
            draw = ImageDraw.Draw(image)
            for n in range(15):
                draw.line([(n * 60, 0), (n * 60 + 220, 500)], fill="#95a5b0", width=8)
            image.save(current / f"GDL-003 microscopy {number}.png")
        (current / "pressure.csv").write_text("time,pressure\n0,12\n1,15\n2,20")
        (legacy / "pressure.csv").write_text("historical,pressure\n0,4")
        document = Document()
        document.add_heading("Compression results", 0)
        document.add_paragraph("Research fixture report")
        document.save(current / "report.docx")
        book = Workbook()
        book.active.append(["pressure", "porosity"])
        book.active.append([20, 71])
        book.save(current / "results.xlsx")
        book.close()
        pdf = PdfWriter()
        pdf.add_blank_page(width=400, height=600)
        pdf.add_blank_page(width=400, height=600)
        pdf.write(current / "report.pdf")
        catalog = NativeIndex(store, locations)
        catalog.refresh()
        repo = ResearchWorkspace(catalog)
        repo.save(
            "sample",
            {
                "name": "Carbon paper sample",
                "material": "Carbon",
                "status": "In testing",
                "notes": "Fixture only; no experimental claims.",
            },
            "GDL-003",
        )
        repo.save(
            "experiment",
            {
                "name": "Optical microscopy",
                "type": "Optical Microscopy",
                "sample_id": "GDL-003",
                "status": "Completed",
                "conditions": [{"name": "Magnification", "value": "20", "unit": "x"}],
                "results": [
                    {
                        "name": "Porosity",
                        "value": "71",
                        "unit": "%",
                        "notes": "Synthetic fixture",
                    }
                ],
            },
            "EXP-003",
        )
        repo.annotate(catalog.query(origin="current")[0], {"experiment_id": "EXP-003"})
    with patch("app.ui.window.QTimer.singleShot"):
        window = HubWindow(store)
    hub = window.research_workspace
    window.resize(1440, 980)
    window.show()
    app.processEvents()
    folder = ROOT / ".test-state/workspace-smoke"
    folder.mkdir(parents=True, exist_ok=True)
    try:
        assert hub.catalog is not None
        for page in (
            "Overview",
            "Samples",
            "Experiments",
            "Images",
            "Data",
            "Reports",
            "Timeline",
            "Files",
            "Legacy",
            "Favorites",
            "Recent",
            "Needs review",
            "Folders",
        ):
            hub.navigate(page)
            wait(app, hub)
            assert window.grab().save(str(folder / (page.replace(" ", "-") + ".png")))
        if "--personal" not in sys.argv:
            hub.show_object("GDL-003")
            hub.object_tabs.setCurrentIndex(1)
            wait(app, hub)
            assert window.grab().save(str(folder / "sample-images.png"))
            hub.show_object("EXP-003")
            hub.object_tabs.setCurrentIndex(4)
            assert window.grab().save(str(folder / "experiment-results.png"))
            for name in ("report.pdf", "pressure.csv", "report.docx"):
                row = next(
                    r
                    for r in hub.catalog.query(origin="current")[0]
                    if r["name"] == name
                )
                hub.navigate("Files")
                wait(app, hub)
                hub.select_file(row)
                wait(app, hub)
                assert window.grab().save(str(folder / ("preview-" + name + ".png")))
            images = hub.catalog.query(family="Images", origin="current")[0]
            hub.file_action("preview", images[0])
            hub.file_action("compare", images[0], images[:2])
            wait(app, hub)
            for i, dialog in enumerate(hub.dialogs):
                assert dialog.grab().save(str(folder / ("viewer-" + str(i) + ".png")))
                dialog.close()
            hub.navigate("Files")
            wait(app, hub)
            window.resize(1024, 700)
            app.processEvents()
            assert window.width() == 1024
            assert window.grab().save(str(folder / "laptop.png"))
        print(
            json.dumps(
                {
                    "version": __import__("app").__version__,
                    "current": hub.catalog.summary()["current"],
                    "legacy": hub.catalog.summary()["legacy"],
                    "screenshots": len(list(folder.glob("*.png"))),
                    "external_apps_launched": False,
                }
            )
        )
    finally:
        wait(app, hub)
        window.close()
        window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        app.processEvents()
        temporary.cleanup()


if __name__ == "__main__":
    run()
