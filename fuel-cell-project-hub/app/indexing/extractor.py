"""Bounded local parsers. No downloads, formula evaluation, OCR or execution."""
import json
import zipfile
from pathlib import Path
from app.services.project_storage import placeholder

DEFAULT_LIMITS = dict(enabled=True, max_file_bytes=16 * 1024 * 1024, max_text_chars=200000,
    max_cells=2000, max_pages=100, max_archive_bytes=32 * 1024 * 1024,
    max_hash_bytes=1024 * 1024, hash_budget=32 * 1024 * 1024)
TEXT = {'.txt', '.md', '.csv', '.json', '.py', '.m', '.xml', '.yaml', '.yml',
        '.jsl', '.ijm', '.ps1', '.bat', '.ini', '.toml', '.tsv', '.r', '.log'}
IMAGES = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.bmp', '.gif'}
SUPPORTED = TEXT | IMAGES | {'.pdf', '.docx', '.xlsx', '.pptx'}


def extract(path, limits=None):
    limits = {**DEFAULT_LIMITS, **(limits or {})}
    path = Path(path)
    result = {'text': '', 'metadata': {}, 'status': 'Metadata only', 'error': ''}
    try:
        info = path.stat()
        if placeholder(info):
            result['status'] = 'Online-only'
            return result
        if not limits['enabled']:
            result['status'] = 'Extraction disabled'
            return result
        if path.suffix.lower() not in SUPPORTED:
            result['status'] = 'Unsupported'
            return result
        if info.st_size > limits['max_file_bytes'] and path.suffix.lower() not in IMAGES:
            result['status'] = 'Partial · file size limit'
            return result
        cap = limits['max_text_chars']
        pieces, used, partial = [], 0, False
        def add(value):
            nonlocal used, partial
            text = str(value or '')
            remaining = cap - used
            if len(text) > remaining:
                partial = True
            value = text[:max(0, remaining)]
            if value:
                pieces.append(value)
                used += len(value) + 1
        extension = path.suffix.lower()
        if extension in TEXT:
            with path.open('rb') as stream:
                raw = stream.read(cap * 4 + 4)
            text = raw.decode('utf-8-sig', errors='replace')
            if '\x00' in text[:4096]:
                # UTF-16 text is common in Windows exports; avoid binary garbage.
                if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
                    text = raw.decode('utf-16', errors='replace')
                else:
                    result['status'] = 'Unsupported binary content'
                    return result
            add(text)
            partial |= len(raw) < info.st_size
        elif extension in IMAGES:
            from PIL import Image
            with Image.open(path) as image:
                result['metadata'] = {'width': image.width, 'height': image.height, 'format': image.format}
                result['status'] = 'Image metadata'
                return result
        elif extension == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            result['metadata'] = {'title': str((reader.metadata or {}).get('/Title', '')),
                                  'page_count': len(reader.pages)}
            for index, page in enumerate(reader.pages):
                if index >= limits['max_pages'] or used >= cap:
                    partial = True
                    break
                contents = page.get_contents()
                if contents and len(contents.get_data()) > limits['max_archive_bytes']:
                    partial = True
                    continue
                add(page.extract_text() or '')
            if not pieces:
                result['status'] = 'OCR required' if not partial else 'Partial · PDF content limit'
                return result
        else:
            # Check expanded OOXML size before constructing document objects.
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > limits['max_archive_bytes']:
                    result['status'] = 'Partial · expanded document size limit'
                    return result
            if extension == '.docx':
                from docx import Document
                document = Document(path)
                result['metadata']['title'] = document.core_properties.title
                for paragraph in document.paragraphs:
                    add(paragraph.text)
                    if used >= cap:
                        partial = True
                        break
                cells = 0
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cells >= limits['max_cells'] or used >= cap:
                                partial = True
                                break
                            add(cell.text)
                            cells += 1
                        if partial:
                            break
                    if partial:
                        break
            elif extension == '.xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                try:
                    result['metadata'].update(workbook_name=path.name, sheet_names=workbook.sheetnames,
                        sheets={s.title: {'rows': s.max_row, 'columns': s.max_column} for s in workbook})
                    cells = 0
                    for sheet in workbook:
                        add(sheet.title)
                        for row in sheet.iter_rows(max_row=limits['max_cells'], max_col=min(sheet.max_column or 1, 100), values_only=True):
                            for value in row:
                                cells += 1
                                if isinstance(value, str) and not value.startswith('='):
                                    add(value)
                                if cells >= limits['max_cells'] or used >= cap:
                                    partial = True
                                    break
                            if partial:
                                break
                        if partial:
                            break
                finally:
                    workbook.close()
            elif extension == '.pptx':
                from pptx import Presentation
                presentation = Presentation(path)
                result['metadata'].update(title=presentation.core_properties.title, slide_count=len(presentation.slides))
                for index, slide in enumerate(presentation.slides):
                    if index >= limits['max_pages'] or used >= cap:
                        partial = True
                        break
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            add(shape.text)
        result.update(text='\n'.join(pieces)[:cap], status='Partial' if partial else 'Content indexed')
    except Exception as exc:
        # File parser errors are local diagnostics, never a reason to abort a scan.
        result['metadata']['parser_error'] = type(exc).__name__
        result.update(status='Extraction failed', error='Unable to parse file; metadata indexed and content skipped (' + type(exc).__name__ + ')')
    return result
