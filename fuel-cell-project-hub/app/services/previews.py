"""Bounded, resident-only preview reads and disposable application-owned thumbnails."""

import csv
import hashlib
import io
import json
import os
import uuid
import zipfile
from itertools import islice
from pathlib import Path
from app.services.project_storage import placeholder

MAX_BYTES = 32 * 1024 * 1024
MAX_PIXELS = 40_000_000


def resident_path(catalog, row):
    path = catalog.safe_path(row)
    if not path.is_file():
        raise ValueError(
            "File currently unavailable. Locate it or reconnect its source."
        )
    info = path.stat()
    if placeholder(info):
        raise ValueError(
            "Online-only file. Open File to download it through OneDrive, then preview again."
        )
    if info.st_size > MAX_BYTES:
        raise ValueError(
            "Preview size limit reached. Open File to view the complete file."
        )
    return path


def check_archive(path):
    with zipfile.ZipFile(path) as archive:
        if sum(x.file_size for x in archive.infolist()) > MAX_BYTES:
            raise ValueError(
                "Preview size limit reached. Open File to view the complete file."
            )


def thumbnail(catalog, row, size=240):
    from PIL import Image, ImageOps

    path = resident_path(catalog, row)
    key = hashlib.sha256(
        json.dumps(
            [row["id"], row.get("signature"), path.stat().st_mtime_ns, size]
        ).encode()
    ).hexdigest()
    folder = Path(catalog.locations.value["cache"]) / "thumbnails"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (key + ".png")
    if target.is_file():
        return str(target)
    with Image.open(path) as image:
        if image.width * image.height > MAX_PIXELS:
            raise ValueError("Image exceeds preview limit. Open File to view it.")
        image.draft("RGB", (size, size))
        image = ImageOps.exif_transpose(image)
        image.thumbnail((size, size))
        temporary = target.with_name(key + "-" + uuid.uuid4().hex + ".tmp")
        try:
            image.convert("RGB").save(temporary, "PNG")
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    return str(target)


def image_path(catalog, row):
    from PIL import Image

    path = resident_path(catalog, row)
    with Image.open(path) as image:
        if image.width * image.height > MAX_PIXELS:
            raise ValueError("Image exceeds viewer limit. Open File to view it.")
    return str(path)


def sheets(catalog, row):
    from openpyxl import load_workbook

    path = resident_path(catalog, row)
    check_archive(path)
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return book.sheetnames
    finally:
        book.close()


def table(catalog, row, sheet=None):
    path = resident_path(catalog, row)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        check_archive(path)
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            page = book[sheet] if sheet else book.active
            return [
                [str(v) if v is not None else "" for v in values]
                for values in page.iter_rows(
                    min_row=1,
                    max_row=250,
                    max_col=min(page.max_column or 1, 100),
                    values_only=True,
                )
            ]
        finally:
            book.close()
    with path.open("rb") as stream:
        data = stream.read(2 * 1024 * 1024)
    decoded = data.decode("utf-8-sig", errors="replace")
    return [
        r[:100]
        for r in islice(
            csv.reader(
                io.StringIO(decoded),
                delimiter="\t" if path.suffix.lower() == ".tsv" else ",",
            ),
            250,
        )
    ]


def text(catalog, row):
    path = resident_path(catalog, row)
    if path.suffix.lower() == ".docx":
        from docx import Document

        check_archive(path)
        document = Document(path)
        paragraphs = [p.text for p in islice(document.paragraphs, 1000)]
        for table in islice(document.tables, 20):
            paragraphs.extend(
                " | ".join(c.text for c in r.cells[:50])
                for r in islice(table.rows, 100)
            )
        return "\n".join(paragraphs)[:100000]
    if path.suffix.lower() in (
        ".txt",
        ".md",
        ".py",
        ".m",
        ".json",
        ".js",
        ".ts",
        ".xml",
        ".yaml",
        ".yml",
        ".log",
        ".bat",
        ".ps1",
        ".cmd",
    ):
        with path.open("rb") as stream:
            data = stream.read(100000)
        if b"\x00" in data:
            raise ValueError(
                "Text preview is unavailable for this encoding. Open File to view it."
            )
        return data.decode("utf-8-sig", errors="replace")
    result = catalog.content_preview(row["id"])
    if result.get("preview"):
        return result["preview"]
    raise ValueError(
        "Preview unavailable for this file type. Open File to use its usual application."
    )
