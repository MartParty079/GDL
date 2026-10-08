"""Shared meeting metadata with user-scoped durable edits and portable file refs."""
import copy
import json
import sqlite3
import uuid
from contextlib import contextmanager
from pathlib import Path, PurePosixPath
from app.services.storage import timestamp, read_json, write_json
from app.services.accounts import AccountError
from app.services.contributions import fetch_pages

MEETING_TYPES = ('Team Meeting', 'Advisor Meeting', 'Lab Meeting', 'Design Review',
                 'Testing Session', 'Presentation Practice', 'Other')
ATTENDANCE = ('Present', 'Absent', 'Remote', 'Excused', 'Partial')


class Meetings:
    def __init__(self, accounts, catalog):
        self.accounts, self.catalog = accounts, catalog
        store = accounts.store
        self.shared = store.shared_index
        self.path = store.local_dir / ('cache/meetings.sqlite3' if self.shared else 'meetings.local.sqlite3')
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.user = accounts.profile['id']
        self.conflicts = []
        with self.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS meeting_cache(user_id TEXT,id TEXT,payload TEXT,pending INTEGER DEFAULT 0,PRIMARY KEY(user_id,id))')
            if self.shared:
                folder = store.project_data('metadata/meetings/' + accounts.install_id)
                rows = [read_json(p, {}) for p in folder.glob('*.json')]
                for entry in sorted(rows, key=lambda r: r.get('updated_at', '')):
                    row = entry.get('record', {})
                    if entry.get('user_id') == self.user and row.get('created_by') and row.get('id'):
                        db.execute('INSERT OR REPLACE INTO meeting_cache VALUES(?,?,?,?)',
                            (self.user, row['id'], json.dumps(row), int(entry.get('pending', True))))

    @contextmanager
    def connect(self):
        db=sqlite3.connect(self.path, timeout=15)
        try:
            with db:
                yield db
        finally:
            db.close()

    def list(self, query=''):
        with self.connect() as db:
            rows = [json.loads(r[0]) for r in db.execute('SELECT payload FROM meeting_cache WHERE user_id=?', (self.user,))]
        return sorted((r for r in rows if query.casefold() in json.dumps(r, ensure_ascii=False).casefold()),
                      key=lambda r: r['meeting_date'], reverse=True)

    def get(self, identity):
        with self.connect() as db:
            row = db.execute('SELECT payload FROM meeting_cache WHERE user_id=? AND id=?', (self.user, identity)).fetchone()
        return json.loads(row[0]) if row else None

    def new(self):
        return dict(id=str(uuid.uuid4()), title='', meeting_date=timestamp()[:10], start_time='', end_time='',
                    location='', meeting_type='Team Meeting', description='', source_type='current',
                    created_by=self.user, updated_by=self.user, revision=0,
                    content=dict(notes='', transcript='', transcript_status='No Transcript',
                                 transcript_history=[], related=[], files={}), attendees=[], actions=[])

    def editable(self, row):
        return row['created_by'] == self.user or self.accounts.profile['role'] == 'admin'

    def save(self, row):
        if not self.editable(row):
            raise ValueError('Only the meeting creator or an administrator can edit this meeting.')
        row = copy.deepcopy(row)
        row['title'] = row['title'].strip()
        if not row['title'] or len(row['title']) > 180:
            raise ValueError('Enter a meeting title of up to 180 characters.')
        if len(json.dumps(row)) > 500000:
            raise ValueError('The meeting is too large. Keep longer transcripts in shared storage.')
        previous = self.get(row['id'])
        row['updated_by'] = self.user
        row['updated_at'] = timestamp()
        row['client_mutation_id'] = str(uuid.uuid4())
        if previous and previous['content'].get('transcript') != row['content'].get('transcript'):
            history = row['content'].setdefault('transcript_history', [])
            history.append(dict(text=previous['content'].get('transcript', ''), at=previous.get('updated_at'), edited_by=previous.get('updated_by')))
            row['content']['transcript_history'] = history[-20:]
        if row['content'].get('transcript') != (previous or {}).get('content',{}).get('transcript',''):
            row['content']['transcript_edited_at'] = row['updated_at']
            row['content']['transcript_edited_by'] = self.user
        if self.shared:
            write_json(self.accounts.store.project_data('metadata/meetings/' + self.accounts.install_id) /
                       (row['client_mutation_id'] + '.json'), {'user_id': self.user, 'updated_at': row['updated_at'], 'record': row, 'pending': True})
        with self.connect() as db:
            db.execute('INSERT INTO meeting_cache VALUES(?,?,?,1) ON CONFLICT(user_id,id) DO UPDATE SET payload=excluded.payload,pending=1',
                       (self.user, row['id'], json.dumps(row)))
        self.accounts.event('MEETING_CREATED' if not previous else 'MEETING_UPDATED', 'meeting', row['id'], row['title'])
        for field, kind in [('notes', 'MEETING_NOTES_EDITED'), ('transcript', 'TRANSCRIPT_EDITED')]:
            if row['content'].get(field) != (previous or {}).get('content', {}).get(field, ''):
                self.accounts.event(kind, 'meeting', row['id'], row['title'])
        previous_actions = {a['id']: a for a in (previous or {}).get('actions', [])}
        for action in row.get('actions', []):
            if action.get('completed') and not previous_actions.get(action['id'], {}).get('completed'):
                self.accounts.event('ACTION_ITEM_COMPLETED', 'meeting', row['id'], row['title'])
        return row

    def sync(self):
        if self.accounts.offline:
            return 'Offline · meeting edits saved locally'
        self.conflicts = []
        with self.connect() as db:
            pending = [(identity, json.loads(payload)) for identity, payload in db.execute(
                'SELECT id,payload FROM meeting_cache WHERE user_id=? AND pending=1', (self.user,))]
        for identity, row in pending:
            try:
                saved = self.accounts.request('POST', '/rest/v1/rpc/hub_save_meeting', {'record': row, 'expected_revision': row['revision']})
            except AccountError:
                raise ValueError('Meeting sync is unavailable. Your edits remain saved locally; reconnect and retry.') from None
            if isinstance(saved, dict) and saved.get('conflict'):
                self.conflicts.append(identity)
                continue
            if self.shared:
                write_json(self.accounts.store.project_data('metadata/meetings/' + self.accounts.install_id) /
                           (str(uuid.uuid4()) + '.json'), {'user_id': self.user, 'updated_at': timestamp(), 'record': saved, 'pending': False})
            with self.connect() as db:
                current = db.execute('SELECT payload FROM meeting_cache WHERE user_id=? AND id=?', (self.user, identity)).fetchone()
                if current and json.loads(current[0]) == row:
                    db.execute('UPDATE meeting_cache SET payload=?,pending=0 WHERE user_id=? AND id=?', (json.dumps(saved), self.user, identity))
                elif current:
                    # Local edits made while this snapshot was sending retain the new base revision.
                    local = json.loads(current[0]); local['revision'] = saved['revision']
                    db.execute('UPDATE meeting_cache SET payload=? WHERE user_id=? AND id=?', (json.dumps(local), self.user, identity))
        rows = fetch_pages(self.accounts, 'meetings')
        attendees = fetch_pages(self.accounts, 'meeting_attendees')
        actions = fetch_pages(self.accounts, 'meeting_actions')
        with self.connect() as db:
            for row in rows:
                row['attendees'] = [r for r in attendees if r['meeting_id'] == row['id']]
                row['actions'] = [r for r in actions if r['meeting_id'] == row['id']]
                db.execute('INSERT INTO meeting_cache VALUES(?,?,?,0) ON CONFLICT(user_id,id) DO UPDATE SET payload=excluded.payload WHERE meeting_cache.pending=0',
                           (self.user, row['id'], json.dumps(row)))
        return ('Conflict · another user saved these meetings. Export your edits, then reload the shared version: ' + ', '.join(self.conflicts)) if self.conflicts else 'Meetings synchronized'

    def discard_pending(self, identity):
        """Explicit UI confirmation required; caller can export the pending draft first."""
        with self.connect() as db:
            db.execute('DELETE FROM meeting_cache WHERE user_id=? AND id=?', (self.user, identity))

    def roster(self):
        return self.accounts.request('POST', '/rest/v1/rpc/hub_team_roster', {}) or []

    def file_ref(self, path):
        path = Path(path).resolve()
        if not self.catalog:
            raise ValueError('Configure research storage before attaching a file.')
        value = self.catalog.locations.value
        for source in [value['active'], *value['legacy'], *value.get('additional', [])]:
            root = Path(source['root_path']).resolve()
            if path.is_relative_to(root):
                if not path.is_file():
                    raise ValueError('Choose an existing shared research file.')
                return dict(source_label=source.get('source_label') or source['name'],
                            relative_path=path.relative_to(root).as_posix(), filename=path.name,
                            size=path.stat().st_size, extension=path.suffix.lower())
        raise ValueError('Choose a file within a configured shared research source.')

    def file_path(self, ref):
        relative = PurePosixPath(ref['relative_path'])
        if relative.is_absolute() or '..' in relative.parts or ':' in str(relative) or '\\' in str(relative):
            raise ValueError('The shared file reference is invalid.')
        value = self.catalog.locations.value
        sources = [s for s in [value['active'], *value['legacy'], *value.get('additional', [])]
                   if (s.get('source_label') or s['name']) == ref['source_label']]
        if len(sources) != 1:
            raise ValueError('Configure a unique matching research source on this computer.')
        root = Path(sources[0]['root_path']).resolve()
        path = (root / str(relative)).resolve()
        if not path.is_relative_to(root):
            raise ValueError('The shared file reference is outside its research source.')
        return path


def import_transcript(path):
    path = Path(path)
    if path.stat().st_size > 15 * 1024 * 1024:
        raise ValueError('Choose a transcript smaller than 15 MB.')
    if path.suffix.lower() in ('.txt', '.md'):
        text = path.read_text(encoding='utf-8-sig')
    elif path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
    elif path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if len(reader.pages) > 200:
            raise ValueError('Choose a transcript with fewer than 200 pages.')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    else:
        raise ValueError('Choose a TXT, MD, DOCX or extractable PDF transcript.')
    if not text.strip():
        raise ValueError('No transcript text was found. Import a text transcript instead.')
    if len(text) > 200000:
        raise ValueError('This transcript is too long for the meeting editor. Link the full shared file.')
    return text


class TranscriptionService:
    """Local model integration seam; normal builds never download/upload recordings."""
    available = False
    status = 'Local transcript generation is not installed. Attach an existing transcript.'

    def generate(self, recording, cancelled=None):
        raise ValueError(self.status)
